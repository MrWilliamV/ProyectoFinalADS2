import csv
from collections import defaultdict
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from functools import wraps

from django.contrib import messages
from django.contrib.auth.decorators import login_required, user_passes_test
from django.core.exceptions import ValidationError
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import F, Sum, Value, Max, DecimalField, Subquery, OuterRef, Case, When, Q
from django.db.models.functions import Coalesce
from django.http import HttpResponse
from django.shortcuts import render, redirect
from django.utils import timezone
from openpyxl import load_workbook

from HealthAndHouse.auth_rol import has_role
from inventory.forms import InventoryMoveForm, InventoryConfigForm
from .models import Inventory, InventoryMovement, ProductLot, LotStock, Product, Location, InventoryConfig, \
    TIPO_MOVIMIENTO_INICIAL
from .queries import get_inventory_filter_data, build_inventory_list_page
from .services import close_period_snapshot, registrar_entrada_compra, registrar_transferencia_lote, \
    registrar_salida_inventario

Q2 = Decimal("0.01")
Q4 = Decimal("0.0001")


def _q2(x: Decimal) -> Decimal:
    return Decimal(x).quantize(Q2, rounding=ROUND_HALF_UP)


def _q4(x: Decimal) -> Decimal:
    return Decimal(x).quantize(Q4, rounding=ROUND_HALF_UP)


def _get_inv(pid, lid):
    if not pid or not lid:
        return None
    return Inventory.objects.filter(product_id=pid, location_id=lid).first()

@login_required()
@has_role("ADMINISTRADOR", "JEFE_ALMACEN")
def inventory_movement_view(request):
    raw_product_id = request.GET.get("product") or request.POST.get("product") or ""
    raw_location_id = request.GET.get("location") or request.POST.get("location") or ""

    def _to_int_or_none(v):
        try:
            return int(v)
        except (TypeError, ValueError):
            return None

    product_id = _to_int_or_none(raw_product_id)
    location_id = _to_int_or_none(raw_location_id)

    inv = _get_inv(product_id, location_id)
    qty_meta = {
        "max": (inv.max_stock if inv and inv.max_stock is not None else None),
        "current": (inv.quantity if inv and inv.quantity is not None else None),
    }

    products = Product.objects.order_by("name")
    locations = Location.objects.order_by("name")

    if request.method == "POST":
        form = InventoryMoveForm(request.POST)
        if form.is_valid():
            cd = form.cleaned_data
            product = cd["product"]
            location = cd["location"]
            quantity = _q2(cd["quantity"])
            unit_cost = _q4(cd.get("unit_price_in") or Decimal("0.0000"))
            movement_type = cd["movement_type"]

            try:
                if movement_type in ["PUR", "ADJIN"]:
                    result = registrar_entrada_compra(
                        user=request.user,
                        product=product,
                        location=location,
                        quantity=quantity,
                        unit_cost=unit_cost,
                        lot_code=cd["lot_code"],
                        expire_date=cd.get("expire_date"),
                        supplier=cd["supplier"],
                        description=cd["descripcion"],
                        movement_type=movement_type,
                    )
                    unit = result["movement"].unidad
                    saldo = _q2(result["inventory"].quantity)
                    messages.success(request, f"Entrada registrada. Saldo actual: {saldo} {unit}.")
                    return redirect("see_inventory")

                elif movement_type in ["SAL", "ADJOUT"]:
                    result = registrar_salida_inventario(
                        user=request.user,
                        product=product,
                        location=location,
                        quantity=quantity,
                        description=cd["descripcion"],
                        movement_type=movement_type,
                    )
                    messages.success(
                        request,
                        f"Salida registrada. Saldo actual: {result['inventory'].quantity} {result['unit']}."
                    )
                    return redirect("see_inventory")

            except ValueError as e:
                form.add_error(None, str(e))
            except Exception as e:
                form.add_error(None, f"Ocurrió un error de sistema: {e}")

        return render(request, "inventory.html", {
            "form": form,
            "qty_meta": qty_meta,
            "products": products,
            "locations": locations,
            "product_id": product_id,
            "location_id": location_id,
        })

    initial = {}
    if product_id is not None:
        initial["product"] = product_id
    if location_id is not None:
        initial["location"] = location_id

    form = InventoryMoveForm(initial=initial)

    return render(request, "inventory.html", {
        "form": form,
        "qty_meta": qty_meta,
        "products": products,
        "locations": locations,
        "product_id": str(product_id) if product_id else "",
        "location_id": str(location_id) if location_id else "",
    })

@login_required
@has_role("ADMINISTRADOR", "JEFE_ALMACEN")
def inventory_list_view(request):
    q = request.GET.get("q", "").strip()
    loc_id = request.GET.get("location")
    sin_ini = request.GET.get("sin_ini") in ("1", "true", "yes")

    period_id = request.GET.get("period")
    if period_id:
        cfg = InventoryConfig.objects.filter(pk=period_id).first()
    else:
        cfg = InventoryConfig.objects.filter(is_active=True).first()

    filter_data = get_inventory_filter_data()
    locations = filter_data["locations"]
    periods = filter_data["periods"]

    if not cfg:
        return render(request, "see_inventory.html", {
            "page_obj": [],
            "locations": locations,
            "current_location": loc_id,
            "search": q,
            "no_active_period": True,
            "sin_ini": sin_ini,
            "periods": periods,
            "current_period_id": None,
        })

    page_obj = build_inventory_list_page(
        cfg=cfg,
        q=q,
        loc_id=loc_id,
        page_number=request.GET.get("page"),
    )

    return render(request, "see_inventory.html", {
        "page_obj": page_obj,
        "locations": locations,
        "current_location": loc_id,
        "search": q,
        "no_active_period": False,
        "sin_ini": sin_ini,
        "periods": periods,
        "current_period_id": cfg.id,
    })

@login_required
@has_role("ADMINISTRADOR", "JEFE_ALMACEN")
def traslado_lote_view(request):
    if request.method == "POST":
        origin_id = request.POST.get("origin")
        target_id = request.POST.get("to_location")
        selected = request.POST.get("selected")
        qty_raw = request.POST.get("quantity")

        if not selected:
            messages.error(request, "Selecciona un lote en la tabla.")
            return redirect("inventory_transfer")

        try:
            lot_id_str, product_id_str = selected.split(":")
            lot_id = int(lot_id_str)
            product_id = int(product_id_str)
        except Exception:
            messages.error(request, "Selección inválida.")
            return redirect("inventory_transfer")

        try:
            result = registrar_transferencia_lote(
                user=request.user,
                origin_id=origin_id,
                target_id=target_id,
                lot_id=lot_id,
                product_id=product_id,
                quantity=qty_raw,
            )
            messages.success(
                request,
                f"Traslado realizado: {result['quantity']} {result['movement_out'].unidad} "
                f"de {result['origin'].name} → {result['target'].name}."
            )
            return redirect("inventory_transfer")

        except ProductLot.DoesNotExist:
            messages.error(request, "El lote seleccionado no existe.")
        except Inventory.DoesNotExist:
            messages.error(request, "No existe inventario en el origen o destino.")
        except LotStock.DoesNotExist:
            messages.error(request, "No hay stock de ese lote en el origen.")
        except ValueError as e:
            messages.error(request, str(e))
        except Exception as e:
            messages.error(request, f"Ocurrió un error: {e}")

        return redirect("inventory_transfer")

    cfg = InventoryConfig.objects.filter(is_active=True).first()

    selected_product = (request.GET.get("product") or "").strip()

    origin_id = request.GET.get("origin")
    if not origin_id:
        first_loc = Location.objects.order_by("id_location").first()
        origin_id = str(first_loc.pk) if first_loc else ""

    locations = Location.objects.all().order_by("id_location")
    products = Product.objects.all().order_by("name")

    if cfg and origin_id:
        ZERO_QTY = Value(0, output_field=DecimalField(max_digits=14, decimal_places=4))

        movs = InventoryMovement.objects.filter(period=cfg, location_id=origin_id)
        agg = (
            movs.values("lot_id", "location_id")
                .annotate(
                    entradas=Coalesce(Sum("entrada_cantidad"), ZERO_QTY),
                    salidas=Coalesce(Sum("salida_cantidad"), ZERO_QTY),
                )
                .annotate(period_qty=F("entradas") - F("salidas"))
        )

        rows = (
            LotStock.objects
            .select_related("lot", "lot__product", "location")
            .filter(location_id=origin_id)
            .annotate(
                product_id=F("lot__product_id"),
                product_name=F("lot__product__name"),
                lot_pk=F("lot_id"),
                lot_code=F("lot__lot_code"),
                expire_date=F("lot__expire_date"),
                period_qty=Coalesce(Subquery(
                    agg.filter(
                        lot_id=OuterRef("lot_id"),
                        location_id=OuterRef("location_id")
                    ).values("period_qty")[:1]
                ), ZERO_QTY),
            )
            .filter(period_qty__gt=0)
            .order_by("product_name", "expire_date", "lot_code")
        )

        if selected_product:
            rows = rows.filter(lot__product_id=selected_product)

        paginator = Paginator(rows, 20)
        page_obj = paginator.get_page(request.GET.get("page"))
    else:
        paginator = Paginator([], 20)
        page_obj = paginator.get_page(1)

    context = {
        "locations": locations,
        "products": products,
        "origin_id": origin_id,
        "selected_product": selected_product,
        "page_obj": page_obj,
        "no_active_period": (cfg is None),
    }
    return render(request, "inventory_transfer.html", context)

def _current_stock_qty(product, location):
    """Stock actual = sum(entradas - salidas) para product+location."""
    agg = InventoryMovement.objects.filter(
        product=product, location=location
    ).aggregate(qty=Coalesce(Sum(F("entrada_cantidad") - F("salida_cantidad")), Value(Decimal("0.00"))))
    return agg["qty"] or Decimal("0.00")


def _stock_limits_for(product, location):
    """
    Obtiene min/max exclusivamente desde Inventory (por producto y ubicación).
    Si no existe registro Inventory, no aplica límites (None, None).
    """
    inv = Inventory.objects.filter(product=product, location=location).only("min_stock", "max_stock").first()
    if not inv:
        return (None, None)

    # normaliza a Decimal si son numéricos
    def _to_dec(x):
        if x is None:
            return None
        return Decimal(str(x))

    try:
        min_s = _to_dec(getattr(inv, "min_stock", None))
    except Exception:
        min_s = None
    try:
        max_s = _to_dec(getattr(inv, "max_stock", None))
    except Exception:
        max_s = None
    return (min_s, max_s)


# Encabezados EXACTOS del import/export de inventario inicial
INITIAL_HEADERS = [
    "codigo_producto",
    "nombre_producto",
    "lote",
    "fecha_vencimiento",
    "cantidad_de_producto",
    "precio",
    "unidad",
]


def _lot_code(lot: ProductLot) -> str:
    for f in ("code", "lot_code", "number", "name"):
        if hasattr(lot, f):
            return str(getattr(lot, f) or "")
    return ""


def _lot_expire(lot: ProductLot):
    for f in ("expire_date", "expiration_date", "due_date", "fecha_vencimiento"):
        if hasattr(lot, f):
            return getattr(lot, f)
    return None

def _product_unit(prod: Product) -> str:
    for f in ("unit", "unidad", "uom", "unidad_medida"):
        if hasattr(prod, f):
            val = getattr(prod, f)
            if val:
                return str(val)
    return "und"

def _avg_cost(prod: Product, location) -> Decimal:
    inv = Inventory.objects.filter(product=prod, location=location).first()
    if inv and getattr(inv, "avg_unit_cost", None) is not None:
        return inv.avg_unit_cost
    return Decimal("0.0000")


def _has_field(model_cls, name: str) -> bool:
    return any(f.name == name for f in model_cls._meta.get_fields())

@login_required()
@has_role("ADMINISTRADOR", "JEFE_ALMACEN")
def inventory_initial_import_view(request):
    from decimal import Decimal

    config = (InventoryConfig.objects
              .filter(is_active=True)
              .select_related("main_location")
              .first())
    if not config:
        messages.info(request, "Configura primero la fecha de corte y la bodega principal.")
        return redirect("inventory_initial_hub")

    template_name = "inventory_initial_import.html"

    # Cerrar periodo contable manualmente
    if request.method == "POST" and "force_close_period" in request.POST:
        try:
            res = close_period_snapshot(request.user, config.pk)
            messages.success(request, f"Periodo cerrado. Snapshot #{res['snapshot_id']} con {res['items']} ítems.")
        except Exception as e:
            messages.error(request, str(e))
        return redirect("inventory_initial_hub")

    # Importar inventario inicial (sin modificaciones)
    if request.method == "POST" and ("import_initial" in request.POST or request.FILES.get("file")):
        file = request.FILES.get("file")
        if not file:
            messages.error(request, "Adjunta un archivo Excel.")
            return render(request, template_name, {"config": config})

        try:
            wb = load_workbook(file, data_only=True)
            ws = wb.active
        except Exception as e:
            messages.error(request, f"Archivo inválido: {e}")
            return render(request, template_name, {"config": config})

        header_row = next(ws.iter_rows(min_row=1, max_row=1))
        headers_exact = [str(c.value).strip() for c in header_row if c.value]
        missing = [h for h in INITIAL_HEADERS if h not in headers_exact]
        if missing:
            messages.error(request, f"Faltan columnas obligatorias: {', '.join(missing)}")
            return render(request, template_name, {"config": config})

        idx = {h: headers_exact.index(h) for h in headers_exact}
        cell = lambda row, key: row[idx[key]].value if key in idx else None

        def q2(x):
            return (x if isinstance(x, Decimal) else Decimal(str(x or "0"))).quantize(Decimal("0.01"))

        def q4(x):
            return (x if isinstance(x, Decimal) else Decimal(str(x or "0"))).quantize(Decimal("0.0000"))

        errors, rows, seen = [], [], set()
        replace = bool(request.POST.get("replace") or request.POST.get("replace_existing"))

        for row in ws.iter_rows(min_row=2):
            cod = cell(row, "codigo_producto")
            nombre = cell(row, "nombre_producto")
            lote = cell(row, "lote")
            venc = cell(row, "fecha_vencimiento")
            qty = cell(row, "cantidad_de_producto")
            cost = cell(row, "precio")
            unidad = cell(row, "unidad")

            if all(x in (None, "", 0) for x in (cod, nombre, lote, venc, qty, cost, unidad)):
                continue

            try:
                prod = Product.objects.get(id_product=int(cod))
            except Exception:
                errors.append(f"Producto inválido: {cod}")
                continue

            try:
                qty = q2(qty)
                cost = q4(cost)
                if qty <= 0:
                    errors.append(f"Cantidad no positiva para producto {cod}")
                    continue
            except Exception:
                errors.append(f"Cantidad o precio inválidos para producto {cod}")
                continue

            if isinstance(venc, datetime):
                venc = venc.date()
            elif isinstance(venc, str) and venc.strip():
                try:
                    venc = datetime.strptime(venc.strip(), "%Y-%m-%d").date()
                except Exception:
                    errors.append(f"Fecha vencimiento inválida para producto {cod}: {venc}")
                    continue
            else:
                venc = None

            key = (prod.pk, str(lote or "").strip())
            if key in seen:
                errors.append(f"Duplicado en archivo: producto {cod} lote {lote}")
                continue
            seen.add(key)

            rows.append({
                "product": prod,
                "lot_code": str(lote or "").strip(),
                "expire_date": venc,
                "qty": qty,
                "unit_cost": cost,
                "unidad": str(unidad or "UND").strip() or "UND",
            })

        if errors:
            for e in errors:
                messages.error(request, e)
            messages.error(request, "La importación fue abortada por errores en el archivo.")
            return render(request, template_name, {"config": config})

        created = 0
        try:
            with transaction.atomic():
                now = timezone.now()
                main_loc = config.main_location
                fecha_dt = datetime.combine(config.fecha_corte, datetime.min.time())
                cfg = config

                if replace:
                    InventoryMovement.objects.filter(
                        period=cfg, location=main_loc, tipo_movimiento="INI",
                        fecha__date=config.fecha_corte,
                    ).delete()
                    LotStock.objects.filter(location=main_loc).delete()
                    Inventory.objects.filter(location=main_loc).update(
                        quantity=Decimal("0.00"),
                        avg_unit_cost=Decimal("0.0000"),
                        updated_at=now,
                    )

                prod_tot_qty = defaultdict(lambda: Decimal("0.00"))
                prod_tot_val = defaultdict(lambda: Decimal("0.0000"))

                for r in rows:
                    prod, lot_code, expire_dt = r["product"], r["lot_code"], r["expire_date"]
                    qty, unit_cost, unidad_txt = r["qty"], r["unit_cost"], r["unidad"]

                    pl_defaults = {"expire_date": expire_dt}
                    if _has_field(ProductLot, "updated_at"):
                        pl_defaults["updated_at"] = now
                    lot, _ = ProductLot.objects.select_for_update().get_or_create(
                        product=prod, lot_code=lot_code, defaults=pl_defaults
                    )
                    if _has_field(ProductLot, "updated_at"):
                        lot.updated_at = now
                        lot.save(update_fields=["expire_date", "updated_at"])
                    else:
                        lot.save(update_fields=["expire_date"])

                    ls_defaults = {"quantity": Decimal("0.00")}
                    if _has_field(LotStock, "updated_at"):
                        ls_defaults["updated_at"] = now
                    ls, _ = LotStock.objects.select_for_update().get_or_create(
                        lot=lot, location=main_loc, defaults=ls_defaults
                    )
                    ls.quantity = q2(qty)
                    if _has_field(LotStock, "updated_at"):
                        ls.updated_at = now
                        ls.save(update_fields=["quantity", "updated_at"])
                    else:
                        ls.save(update_fields=["quantity"])

                    prod_tot_qty[(prod.pk, main_loc.pk)] += qty
                    prod_tot_val[(prod.pk, main_loc.pk)] += q4(qty * unit_cost)

                    InventoryMovement.objects.create(
                        product=prod, location=main_loc, lot=lot,
                        fecha=fecha_dt, tipo_movimiento="INI",
                        descripcion="Inventario Inicial",
                        valor_unitario=unit_cost,
                        entrada_cantidad=qty, entrada_valor=q4(qty * unit_cost),
                        salida_cantidad=Decimal("0.00"), salida_valor=Decimal("0.0000"),
                        saldo_cantidad=qty, saldo_valor=q4(qty * unit_cost),
                        proveedor=None, unidad=unidad_txt,
                        created_by=(request.user if request.user.is_authenticated else None),
                        period=cfg,
                    )
                    created += 1

                for (prod_id, loc_id), tot_qty in prod_tot_qty.items():
                    tot_val = prod_tot_val[(prod_id, loc_id)]
                    avg = q4(tot_val / tot_qty) if tot_qty > 0 else Decimal("0.0000")
                    inv, created = Inventory.objects.select_for_update().get_or_create(
                        product_id=prod_id,
                        location_id=loc_id,
                        defaults={
                            "quantity": Decimal("0.00"),
                            "avg_unit_cost": Decimal("0.0000"),
                            "updated_at": now,
                            "min_stock": Decimal("0.00"),
                            "max_stock": Decimal("0.00"),
                        },
                    )

                    inv.quantity = q2(tot_qty)
                    inv.avg_unit_cost = avg
                    if _has_field(Inventory, "updated_at"):
                        inv.updated_at = now
                        inv.save(update_fields=["quantity", "avg_unit_cost", "updated_at"])
                    else:
                        inv.save(update_fields=["quantity", "avg_unit_cost"])

            messages.success(
                request,
                f"Inventario inicial importado correctamente ({created} filas) en {config.main_location.name}."
            )
            return redirect("initial_inventory")

        except ValidationError as ve:
            messages.error(request, f"Error de validación: {ve}")
        except Exception as ex:
            messages.error(request, f"Ocurrió un error al importar: {ex}")

    return render(request, template_name, {"config": config})

def require_active_inventory_config(view_func):
    """Protege vistas: exige InventoryConfig activa, si no, redirige al setup."""

    @wraps(view_func)
    def _wrapped(request, *args, **kwargs):
        if not InventoryConfig.objects.filter(is_active=True).exists():
            messages.info(request, "Configura primero la fecha de corte y la bodega principal.")
            return redirect("inventory_initial_hub")
        return view_func(request, *args, **kwargs)

    return _wrapped

@login_required
@has_role("ADMINISTRADOR")
def inventory_initial_hub(request):
    """
    Hub:
      - Si NO hay InventoryConfig activa, muestra el formulario para crearla.
      - Si SÍ hay, muestra la interfaz para importar inventario inicial.
      - Al crear un nuevo periodo, aplica automáticamente el snapshot del periodo anterior:
          1) Busca snapshot del periodo anterior
          2) Si no existe, lo construye desde movimientos del periodo anterior
          3) Lo aplica al periodo recién creado
    """
    from .services import get_previous_period, aplicar_snapshot_al_periodo_activo, construir_snapshot_desde_movimientos_periodo
    from .models import InventoryPeriodSnapshot

    active = InventoryConfig.objects.filter(is_active=True).select_related("main_location").first()

    if not active:
        if request.method == "POST":
            form = InventoryConfigForm(request.POST)
            if form.is_valid():
                # 1) CREAR el nuevo periodo (activo)
                cfg = form.save()

                # 2) BUSCAR PERIODO ANTERIOR (ya cerrado)
                prev_cfg = get_previous_period(cfg)
                print(f"prev_cfg: {prev_cfg}")
                if prev_cfg:
                    # 2a) Buscar snapshot real
                    snap = (
                        InventoryPeriodSnapshot.objects
                        .filter(periodo=prev_cfg)
                        .order_by("-cutoff_ts")
                        .first()
                    )

                    # 2b) Si no hay snapshot, CONSTRUIRLO desde movimientos del periodo anterior
                    if not snap:
                        snap = construir_snapshot_desde_movimientos_periodo(request.user, prev_cfg)

                    if snap:
                        # 3) APLICAR al periodo recién creado (cfg)
                        try:
                            res = aplicar_snapshot_al_periodo_activo(request.user, snap.id)
                            messages.success(
                                request,
                                f"Periodo creado y snapshot #{snap.id} aplicado automáticamente: "
                                f"{res['ini_movements']} INI, {res['lot_items']} lotes, {res['products']} productos."
                            )
                        except Exception as e:
                            messages.error(request, f"Se creó el periodo, pero no se pudo aplicar el snapshot: {e}")
                    else:
                        messages.info(request, "Periodo creado. El anterior no tenía saldos para snapshot.")
                else:
                    messages.info(request, "Periodo creado. No existe un periodo anterior cerrado para aplicar.")

                return redirect("inventory_initial_hub")
            else:
                messages.error(request, "Por favor corrige los campos marcados.")
        else:
            form = InventoryConfigForm()
        return render(request, "inventory_initial_setup.html", {"form": form})

    # Ya hay activo → mostrar interfaz
    return render(request, "inventory_initial_import.html", {"config": active})

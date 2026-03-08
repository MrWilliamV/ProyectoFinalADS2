
from collections import defaultdict
from datetime import datetime
from functools import wraps

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import ValidationError
from django.db import transaction

from django.shortcuts import render, redirect
from django.utils import timezone
from openpyxl import load_workbook

from HealthAndHouse.auth_rol import has_role
from inventory.forms import  InventoryConfigForm
from inventory.models import InventoryMovement, Inventory, ProductLot, InventoryConfig, LotStock, \
    InventoryPeriodSnapshot
from inventory.queries import INITIAL_HEADERS, _has_field
from inventory.services import close_period_snapshot, get_previous_period, construir_snapshot_desde_movimientos_periodo, \
    aplicar_snapshot_al_periodo_activo
from product.models import Product

@login_required()
@has_role("ADMINISTRADOR", "JEFE_ALMACEN")
def inventory_initial_import_view(request):
    from decimal import Decimal

    config = (InventoryConfig.objects
              .filter(is_active=True)
              .select_related("main_location")
              .first())
    if not config:
        messages.info(request, "Configura primero la fecha de corte y la bodega principal.")
        return redirect("inventory_initial_hub")

    template_name = "inventory_initial_import.html"

    # Cerrar periodo contable manualmente
    if request.method == "POST" and "force_close_period" in request.POST:
        try:
            res = close_period_snapshot(request.user, config.pk)
            messages.success(request, f"Periodo cerrado. Snapshot #{res['snapshot_id']} con {res['items']} ítems.")
        except Exception as e:
            messages.error(request, str(e))
        return redirect("inventory_initial_hub")

    # Importar inventario inicial (sin modificaciones)
    if request.method == "POST" and ("import_initial" in request.POST or request.FILES.get("file")):
        file = request.FILES.get("file")
        if not file:
            messages.error(request, "Adjunta un archivo Excel.")
            return render(request, template_name, {"config": config})

        try:
            wb = load_workbook(file, data_only=True)
            ws = wb.active
        except Exception as e:
            messages.error(request, f"Archivo inválido: {e}")
            return render(request, template_name, {"config": config})

        header_row = next(ws.iter_rows(min_row=1, max_row=1))
        headers_exact = [str(c.value).strip() for c in header_row if c.value]
        missing = [h for h in INITIAL_HEADERS if h not in headers_exact]
        if missing:
            messages.error(request, f"Faltan columnas obligatorias: {', '.join(missing)}")
            return render(request, template_name, {"config": config})

        idx = {h: headers_exact.index(h) for h in headers_exact}
        cell = lambda row, key: row[idx[key]].value if key in idx else None

        def q2(x):
            return (x if isinstance(x, Decimal) else Decimal(str(x or "0"))).quantize(Decimal("0.01"))

        def q4(x):
            return (x if isinstance(x, Decimal) else Decimal(str(x or "0"))).quantize(Decimal("0.0000"))

        errors, rows, seen = [], [], set()
        replace = bool(request.POST.get("replace") or request.POST.get("replace_existing"))

        for row in ws.iter_rows(min_row=2):
            cod = cell(row, "codigo_producto")
            nombre = cell(row, "nombre_producto")
            lote = cell(row, "lote")
            venc = cell(row, "fecha_vencimiento")
            qty = cell(row, "cantidad_de_producto")
            cost = cell(row, "precio")
            unidad = cell(row, "unidad")

            if all(x in (None, "", 0) for x in (cod, nombre, lote, venc, qty, cost, unidad)):
                continue

            try:
                prod = Product.objects.get(id_product=int(cod))
            except Exception:
                errors.append(f"Producto inválido: {cod}")
                continue

            try:
                qty = q2(qty)
                cost = q4(cost)
                if qty <= 0:
                    errors.append(f"Cantidad no positiva para producto {cod}")
                    continue
            except Exception:
                errors.append(f"Cantidad o precio inválidos para producto {cod}")
                continue

            if isinstance(venc, datetime):
                venc = venc.date()
            elif isinstance(venc, str) and venc.strip():
                try:
                    venc = datetime.strptime(venc.strip(), "%Y-%m-%d").date()
                except Exception:
                    errors.append(f"Fecha vencimiento inválida para producto {cod}: {venc}")
                    continue
            else:
                venc = None

            key = (prod.pk, str(lote or "").strip())
            if key in seen:
                errors.append(f"Duplicado en archivo: producto {cod} lote {lote}")
                continue
            seen.add(key)

            rows.append({
                "product": prod,
                "lot_code": str(lote or "").strip(),
                "expire_date": venc,
                "qty": qty,
                "unit_cost": cost,
                "unidad": str(unidad or "UND").strip() or "UND",
            })

        if errors:
            for e in errors:
                messages.error(request, e)
            messages.error(request, "La importación fue abortada por errores en el archivo.")
            return render(request, template_name, {"config": config})

        created = 0
        try:
            with transaction.atomic():
                now = timezone.now()
                main_loc = config.main_location
                fecha_dt = datetime.combine(config.fecha_corte, datetime.min.time())
                cfg = config

                if replace:
                    InventoryMovement.objects.filter(
                        period=cfg, location=main_loc, tipo_movimiento="INI",
                        fecha__date=config.fecha_corte,
                    ).delete()
                    LotStock.objects.filter(location=main_loc).delete()
                    Inventory.objects.filter(location=main_loc).update(
                        quantity=Decimal("0.00"),
                        avg_unit_cost=Decimal("0.0000"),
                        updated_at=now,
                    )

                prod_tot_qty = defaultdict(lambda: Decimal("0.00"))
                prod_tot_val = defaultdict(lambda: Decimal("0.0000"))

                for r in rows:
                    prod, lot_code, expire_dt = r["product"], r["lot_code"], r["expire_date"]
                    qty, unit_cost, unidad_txt = r["qty"], r["unit_cost"], r["unidad"]

                    pl_defaults = {"expire_date": expire_dt}
                    if _has_field(ProductLot, "updated_at"):
                        pl_defaults["updated_at"] = now
                    lot, _ = ProductLot.objects.select_for_update().get_or_create(
                        product=prod, lot_code=lot_code, defaults=pl_defaults
                    )
                    if _has_field(ProductLot, "updated_at"):
                        lot.updated_at = now
                        lot.save(update_fields=["expire_date", "updated_at"])
                    else:
                        lot.save(update_fields=["expire_date"])

                    ls_defaults = {"quantity": Decimal("0.00")}
                    if _has_field(LotStock, "updated_at"):
                        ls_defaults["updated_at"] = now
                    ls, _ = LotStock.objects.select_for_update().get_or_create(
                        lot=lot, location=main_loc, defaults=ls_defaults
                    )
                    ls.quantity = q2(qty)
                    if _has_field(LotStock, "updated_at"):
                        ls.updated_at = now
                        ls.save(update_fields=["quantity", "updated_at"])
                    else:
                        ls.save(update_fields=["quantity"])

                    prod_tot_qty[(prod.pk, main_loc.pk)] += qty
                    prod_tot_val[(prod.pk, main_loc.pk)] += q4(qty * unit_cost)

                    InventoryMovement.objects.create(
                        product=prod, location=main_loc, lot=lot,
                        fecha=fecha_dt, tipo_movimiento="INI",
                        descripcion="Inventario Inicial",
                        valor_unitario=unit_cost,
                        entrada_cantidad=qty, entrada_valor=q4(qty * unit_cost),
                        salida_cantidad=Decimal("0.00"), salida_valor=Decimal("0.0000"),
                        saldo_cantidad=qty, saldo_valor=q4(qty * unit_cost),
                        proveedor=None, unidad=unidad_txt,
                        created_by=(request.user if request.user.is_authenticated else None),
                        period=cfg,
                    )
                    created += 1

                for (prod_id, loc_id), tot_qty in prod_tot_qty.items():
                    tot_val = prod_tot_val[(prod_id, loc_id)]
                    avg = q4(tot_val / tot_qty) if tot_qty > 0 else Decimal("0.0000")
                    inv, created = Inventory.objects.select_for_update().get_or_create(
                        product_id=prod_id,
                        location_id=loc_id,
                        defaults={
                            "quantity": Decimal("0.00"),
                            "avg_unit_cost": Decimal("0.0000"),
                            "updated_at": now,
                            "min_stock": Decimal("0.00"),
                            "max_stock": Decimal("0.00"),
                        },
                    )

                    inv.quantity = q2(tot_qty)
                    inv.avg_unit_cost = avg
                    if _has_field(Inventory, "updated_at"):
                        inv.updated_at = now
                        inv.save(update_fields=["quantity", "avg_unit_cost", "updated_at"])
                    else:
                        inv.save(update_fields=["quantity", "avg_unit_cost"])

            messages.success(
                request,
                f"Inventario inicial importado correctamente ({created} filas) en {config.main_location.name}."
            )
            return redirect("initial_inventory")

        except ValidationError as ve:
            messages.error(request, f"Error de validación: {ve}")
        except Exception as ex:
            messages.error(request, f"Ocurrió un error al importar: {ex}")

    return render(request, template_name, {"config": config})

@login_required
@has_role("ADMINISTRADOR")
def inventory_initial_hub(request):
    """
    Hub:
      - Si NO hay InventoryConfig activa, muestra el formulario para crearla.
      - Si SÍ hay, muestra la interfaz para importar inventario inicial.
      - Al crear un nuevo periodo, aplica automáticamente el snapshot del periodo anterior:
          1) Busca snapshot del periodo anterior
          2) Si no existe, lo construye desde movimientos del periodo anterior
          3) Lo aplica al periodo recién creado
    """

    active = InventoryConfig.objects.filter(is_active=True).select_related("main_location").first()

    if not active:
        if request.method == "POST":
            form = InventoryConfigForm(request.POST)
            if form.is_valid():
                # 1) CREAR el nuevo periodo (activo)
                cfg = form.save()

                # 2) BUSCAR PERIODO ANTERIOR (ya cerrado)
                prev_cfg = get_previous_period(cfg)
                print(f"prev_cfg: {prev_cfg}")
                if prev_cfg:
                    # 2a) Buscar snapshot real
                    snap = (
                        InventoryPeriodSnapshot.objects
                        .filter(periodo=prev_cfg)
                        .order_by("-cutoff_ts")
                        .first()
                    )

                    # 2b) Si no hay snapshot, CONSTRUIRLO desde movimientos del periodo anterior
                    if not snap:
                        snap = construir_snapshot_desde_movimientos_periodo(request.user, prev_cfg)

                    if snap:
                        # 3) APLICAR al periodo recién creado (cfg)
                        try:
                            res = aplicar_snapshot_al_periodo_activo(request.user, snap.id)
                            messages.success(
                                request,
                                f"Periodo creado y snapshot #{snap.id} aplicado automáticamente: "
                                f"{res['ini_movements']} INI, {res['lot_items']} lotes, {res['products']} productos."
                            )
                        except Exception as e:
                            messages.error(request, f"Se creó el periodo, pero no se pudo aplicar el snapshot: {e}")
                    else:
                        messages.info(request, "Periodo creado. El anterior no tenía saldos para snapshot.")
                else:
                    messages.info(request, "Periodo creado. No existe un periodo anterior cerrado para aplicar.")

                return redirect("inventory_initial_hub")
            else:
                messages.error(request, "Por favor corrige los campos marcados.")
        else:
            form = InventoryConfigForm()
        return render(request, "inventory_initial_setup.html", {"form": form})

    # Ya hay activo → mostrar interfaz
    return render(request, "inventory_initial_import.html", {"config": active})

def require_active_inventory_config(view_func):
    """Protege vistas: exige InventoryConfig activa, si no, redirige al setup."""

    @wraps(view_func)
    def _wrapped(request, *args, **kwargs):
        if not InventoryConfig.objects.filter(is_active=True).exists():
            messages.info(request, "Configura primero la fecha de corte y la bodega principal.")
            return redirect("inventory_initial_hub")
        return view_func(request, *args, **kwargs)

    return _wrapped
